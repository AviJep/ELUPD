from __future__ import annotations

from pathlib import Path
from typing import Any

from django.core.management.base import BaseCommand, CommandError
from openpyxl import load_workbook

from core.models import AnnexClupStatus, AnnexPdpfpStatus, DefinitionStatus


def _normalize_header(value: Any) -> str:
    text = str(value or "").strip().lower()
    return "".join(char for char in text if char.isalnum())


def _as_string(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).replace("\xa0", " ").strip()
    return text if text else None


def _as_int(value: Any) -> int | None:
    if value is None or value == "":
        return None
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


def _as_year(value: Any) -> int | None:
    parsed = _as_int(value)
    if parsed is None or parsed <= 1900:
        return None
    return parsed


def _as_bool(value: Any) -> bool:
    if value in (None, ""):
        return False
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value != 0
    lowered = str(value).strip().lower()
    return lowered in {"1", "true", "yes", "y", "x", "done"}


class Command(BaseCommand):
    help = "Import Annex CLUP/PDPFP and Definition Status data from the NIR workbook."

    def add_arguments(self, parser):
        parser.add_argument(
            "--file",
            type=str,
            default=None,
            help="Path to workbook (.xlsx). Defaults to frontend/data/02_NIR_2026_Accomplishment.xlsx",
        )

    def handle(self, *args, **options):
        backend_dir = Path(__file__).resolve().parents[3]
        project_dir = backend_dir.parent
        default_file = project_dir / "frontend" / "data" / "02_NIR_2026_Accomplishment.xlsx"
        workbook_path = Path(options["file"]).resolve() if options["file"] else default_file

        if not workbook_path.exists():
            raise CommandError(f"Workbook not found: {workbook_path}")

        workbook = load_workbook(workbook_path, data_only=True)

        clup_count = self._import_clup(workbook)
        pdpfp_count = self._import_pdpfp(workbook)
        definition_count = self._import_definition(workbook)

        self.stdout.write(self.style.SUCCESS("Annex data import complete."))
        self.stdout.write(f"CLUP rows imported: {clup_count}")
        self.stdout.write(f"PDPFP rows imported: {pdpfp_count}")
        self.stdout.write(f"Definition rows imported: {definition_count}")

    def _build_header_map(self, sheet) -> dict[str, int]:
        headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        return {_normalize_header(value): index for index, value in enumerate(headers)}

    def _find_header_row(self, sheet, expected_headers: list[str]) -> tuple[int, dict[str, int]]:
        normalized_expected = {_normalize_header(item) for item in expected_headers}
        best_row_index = 1
        best_row_map: dict[str, int] = {}
        best_score = -1

        for row_index in range(1, min(sheet.max_row, 30) + 1):
            row_values = next(sheet.iter_rows(min_row=row_index, max_row=row_index, values_only=True))
            row_map = {
                _normalize_header(value): index
                for index, value in enumerate(row_values)
                if _normalize_header(value)
            }
            score = len(normalized_expected.intersection(row_map.keys()))
            if score > best_score:
                best_score = score
                best_row_index = row_index
                best_row_map = row_map

        return best_row_index, best_row_map

    def _cell(self, row: tuple[Any, ...], header_map: dict[str, int], *aliases: str) -> Any:
        normalized_aliases = [_normalize_header(alias) for alias in aliases]

        for key in normalized_aliases:
            if key in header_map:
                return row[header_map[key]]

        for alias in aliases:
            key = _normalize_header(alias)
            for header_key, index in header_map.items():
                if key and (key in header_key or header_key in key):
                    return row[index]
        return None

    def _import_clup(self, workbook) -> int:
        sheet_name = "Annex 3.1_CLUP Status"
        if sheet_name not in workbook.sheetnames:
            raise CommandError(f"Missing sheet: {sheet_name}")

        sheet = workbook[sheet_name]
        header_row, header_map = self._find_header_row(
            sheet,
            [
                "City/Municipality",
                "Province",
                "Start Year",
                "End Year",
                "Resolution Number of the Latest Plan",
                "CLUP Status",
                "Pre-phase",
                "Phase 1",
                "Phase 2",
                "Phase 3",
                "Phase 4",
                "Phase 5",
                "Preparation Status",
            ],
        )

        AnnexClupStatus.objects.all().delete()

        rows: list[AnnexClupStatus] = []
        for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
            city = _as_string(self._cell(row, header_map, "City/Municipality", "City Municipality"))
            province = _as_string(self._cell(row, header_map, "Province"))
            if not city or not province:
                continue

            rows.append(
                AnnexClupStatus(
                    cityMunicipality=city,
                    province=province,
                    planningStartYear=_as_year(self._cell(row, header_map, "Planning Start Year", "Start Year")),
                    planningEndYear=_as_year(self._cell(row, header_map, "Planning End Year", "End Year")),
                    resolutionNumber=_as_string(self._cell(row, header_map, "Resolution Number", "SP Resolution Number")),
                    clupStatus=_as_string(self._cell(row, header_map, "CLUP Status", "Status")),
                    prePhase=_as_bool(self._cell(row, header_map, "Pre-Phase", "Pre Phase", "Prephase")),
                    phase1=_as_bool(self._cell(row, header_map, "Phase 1", "Phase1")),
                    phase2=_as_bool(self._cell(row, header_map, "Phase 2", "Phase2")),
                    phase3=_as_bool(self._cell(row, header_map, "Phase 3", "Phase3")),
                    phase4=_as_bool(self._cell(row, header_map, "Phase 4", "Phase4")),
                    phase5=_as_bool(self._cell(row, header_map, "Phase 5", "Phase5")),
                    currentProgress=_as_string(self._cell(row, header_map, "Current Progress", "Current Phase", "Preparation Status")) or "",
                )
            )

        AnnexClupStatus.objects.bulk_create(rows, batch_size=500)
        return len(rows)

    def _import_pdpfp(self, workbook) -> int:
        sheet_name = "Annex 3.2_PDPFP Status"
        if sheet_name not in workbook.sheetnames:
            raise CommandError(f"Missing sheet: {sheet_name}")

        sheet = workbook[sheet_name]
        header_row, header_map = self._find_header_row(
            sheet,
            [
                "Region",
                "Province",
                "Income Classification",
                "Version of PPFP",
                "Start Year",
                "End Year",
                "Resolution Approving Plan",
                "Year Approved",
                "Year Adopted",
                "PDPFP/PPFP Status",
                "Technical Assistance Provided",
                "Support from Other Institutions",
                "LGU with Local Shelter Plan",
                "If with Support from other institutions, Name/s of Institutions",
                "Remarks",
            ],
        )

        AnnexPdpfpStatus.objects.all().delete()

        rows: list[AnnexPdpfpStatus] = []
        for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
            province = _as_string(self._cell(row, header_map, "Province"))
            region = _as_string(self._cell(row, header_map, "Region"))
            status = _as_string(self._cell(row, header_map, "Status", "Latest Status"))
            if not province or not region or not status:
                continue

            rows.append(
                AnnexPdpfpStatus(
                    region=region,
                    province=province,
                    incomeClassification=_as_string(self._cell(row, header_map, "Income Classification", "Income Class")),
                    version=_as_string(self._cell(row, header_map, "Version")),
                    startYear=_as_year(self._cell(row, header_map, "Start Year")),
                    endYear=_as_year(self._cell(row, header_map, "End Year")),
                    resolutionApprovingPlan=_as_string(self._cell(row, header_map, "Resolution Approving Plan", "Resolution Number")),
                    yearApproved=_as_year(self._cell(row, header_map, "Year Approved")),
                    yearAdopted=_as_year(self._cell(row, header_map, "Year Adopted")),
                    status=status,
                    technicalAssistance=_as_string(self._cell(row, header_map, "Technical Assistance")),
                    supportFromOtherInstitutions=_as_string(self._cell(row, header_map, "Support from Other Institutions")),
                    withLocalShelterPlan=_as_string(self._cell(row, header_map, "With Local Shelter Plan")),
                    institutions=_as_string(self._cell(row, header_map, "Institutions")),
                    remarks=_as_string(self._cell(row, header_map, "Remarks")),
                )
            )

        AnnexPdpfpStatus.objects.bulk_create(rows, batch_size=200)
        return len(rows)

    def _import_definition(self, workbook) -> int:
        sheet_name = "Definition_Status"
        if sheet_name not in workbook.sheetnames:
            raise CommandError(f"Missing sheet: {sheet_name}")

        sheet = workbook[sheet_name]
        header_row, header_map = self._find_header_row(
            sheet,
            ["Class", "Subclass", "Definition", "Means of Verification"],
        )

        DefinitionStatus.objects.all().delete()

        rows: list[DefinitionStatus] = []
        for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
            class_name = _as_string(self._cell(row, header_map, "Class", "Main Class"))
            subclass = _as_string(self._cell(row, header_map, "Subclass", "Sub-class"))
            definition = _as_string(self._cell(row, header_map, "Definition"))
            means = _as_string(self._cell(row, header_map, "Means of Verification", "MOV"))
            if not class_name and not subclass and not definition and not means:
                continue

            rows.append(
                DefinitionStatus(
                    className=class_name,
                    subclass=subclass,
                    definition=definition,
                    meansOfVerification=means,
                )
            )

        DefinitionStatus.objects.bulk_create(rows, batch_size=200)
        return len(rows)
